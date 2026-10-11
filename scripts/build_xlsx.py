#!/usr/bin/env python3
"""Build the Quiz-Gen exam workbook (.xlsx) from a questions JSON file.

The "Questions" sheet is byte-for-byte the same schema as NotesLM's
quiz-editor _exportExcel():  # | Learning Outcome | Topic | QUESTION |
Type of Question | Taxonomy | Choice A-D | Answer | Correct answer in Words |
Timer (Seconds). Column widths match too.

Timers are recomputed with the NotesLM rules when a question has no timer, so
an agent that omits them still gets correct values:
    Multiple Choice / True-False : max(20, ceil(words * 1.5 / 5) * 5)
                                   (words = question + all choice texts)
    Identification               : Bloom's taxonomy seconds (default 80)
    Problem Solving              : min(600, max(40, 40 * steps))
    Short Answer                 : 900

Input JSON may be either a bare list of question objects, or an object with a
"questions" array plus optional "subjectCode", "description", "examPeriod",
"filename". Question objects use the NotesLM schema: learningOutcome, topic,
question, type, taxonomy, choices {A..D}, answer, answerText, timer.

Usage:
    python build_xlsx.py questions.json
    python build_xlsx.py questions.json -o FM100_PRELIM.xlsx
    python build_xlsx.py questions.json --answer-key
"""
import argparse
import json
import math
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADERS = ["#", "Learning Outcome", "Topic", "QUESTION", "Type of Question",
           "Taxonomy", "Choice A", "Choice B", "Choice C", "Choice D",
           "Answer", "Correct answer in Words", "Timer (Seconds)"]
WIDTHS = [4, 20, 16, 40, 14, 12, 30, 30, 30, 30, 6, 30, 12]
MCQ_TYPES = {"Multiple Choice", "True/False"}
TAXONOMY_SECONDS = {"Remembering": 40, "Understanding": 80, "Applying": 80,
                    "Analyzing": 120, "Evaluating": 120, "Creating": 120}
KEY_HEADERS = ["#", "Question", "Answer", "Answer Text", "Topic", "Taxonomy"]
KEY_WIDTHS = [4, 60, 8, 40, 20, 14]


def word_count_timer(question, choices):
    text = (question or "") + " " + " ".join(str(v) for v in (choices or {}).values())
    words = len([w for w in re.split(r"\s+", text) if w])
    return max(20, math.ceil(words * 1.5 / 5) * 5)


def parse_steps(answer_text):
    if not answer_text:
        return []
    normalized = str(answer_text).replace("->", "→")
    return [s.strip() for s in normalized.split("→") if s.strip()]


def compute_timer(q):
    qtype = q.get("type") or "Multiple Choice"
    if qtype in MCQ_TYPES:
        return word_count_timer(q.get("question"), q.get("choices"))
    if qtype == "Identification":
        return TAXONOMY_SECONDS.get(q.get("taxonomy"), 80)
    if qtype == "Problem Solving":
        steps = max(1, len(parse_steps(q.get("answerText"))))
        return min(600, max(40, 40 * steps))
    if qtype == "Short Answer":
        return 900
    return 30


def row_for(q, i, force_timer=False):
    qtype = q.get("type") or "Multiple Choice"
    mcq = qtype in MCQ_TYPES
    choices = q.get("choices") or {}
    timer = q.get("timer")
    if force_timer or not isinstance(timer, (int, float)) or timer <= 0:
        timer = compute_timer(q)
    return [
        i + 1,
        q.get("learningOutcome", ""),
        q.get("topic", ""),
        q.get("question", ""),
        qtype,
        q.get("taxonomy", "Understanding"),
        choices.get("A", "") if mcq else "",
        choices.get("B", "") if mcq else "",
        ("" if qtype == "True/False" else choices.get("C", "")) if mcq else "",
        ("" if qtype == "True/False" else choices.get("D", "")) if mcq else "",
        q.get("answer", "") if mcq else "",
        q.get("answerText", ""),
        int(timer),
    ]


def load(path):
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    if isinstance(data, list):
        return {"questions": data}, data
    if isinstance(data, dict) and isinstance(data.get("questions"), list):
        return data, data["questions"]
    raise SystemExit("questions JSON must be a list or an object with a 'questions' array")


def build(data, questions, out_path, answer_key=False, force_timer=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(HEADERS)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = bold
        cell.fill = fill
    for i, q in enumerate(questions):
        ws.append(row_for(q, i, force_timer=force_timer))
    for col, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if answer_key:
        ks = wb.create_sheet("Answer Key")
        ks.append(KEY_HEADERS)
        for cell in ks[1]:
            cell.font = bold
        for i, q in enumerate(questions):
            answer = q.get("answer") or q.get("answerText") or ""
            ks.append([i + 1, q.get("question", ""), q.get("answer", ""),
                       q.get("answerText", ""), q.get("topic", ""), q.get("taxonomy", "")])
        for col, width in enumerate(KEY_WIDTHS, start=1):
            ks.column_dimensions[ks.cell(row=1, column=col).column_letter].width = width

    wb.save(out_path)
    return out_path


def default_name(data):
    if data.get("filename"):
        return data["filename"]
    code = str(data.get("subjectCode", "")).strip()
    period = str(data.get("examPeriod", "")).strip()
    if code or period:
        return f"{code}_{period}".strip("_") + ".xlsx"
    return "Quiz.xlsx"


def main():
    ap = argparse.ArgumentParser(description="Build the Quiz-Gen exam .xlsx from questions JSON.")
    ap.add_argument("questions_json")
    ap.add_argument("-o", "--out")
    ap.add_argument("--answer-key", action="store_true", help="add a second 'Answer Key' sheet")
    ap.add_argument("--recompute-timers", action="store_true",
                    help="override any provided timer values with the computed ones")
    args = ap.parse_args()

    data, questions = load(args.questions_json)
    if not questions:
        raise SystemExit("no questions found in JSON")
    out = args.out or default_name(data)
    build(data, questions, out, answer_key=args.answer_key, force_timer=args.recompute_timers)
    print(f"Wrote {len(questions)} questions -> {out}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
